"""
Real-World Acceptance Test: Autonomous Presentation Generation.
Tests PowerPoint (.pptx) creation using python-pptx, verifies file integrity,
slide structures, pitch scripts, and judge Q&A bundles.
"""

from pathlib import Path
import pytest
from pptx import Presentation
from agents.presentation.agent import PresentationAgent
from agents.presentation.models import PresentationMode
from core.artifacts.manager import ArtifactManager


def test_real_presentation_deck_generation(tmp_path: Path):
    """Verify PresentationAgent creates valid .pptx files and rich bundles."""
    art_mgr = ArtifactManager(root_dir=tmp_path / "artifacts")
    agent = PresentationAgent(artifact_manager=art_mgr)
    
    deck = agent.build_deck(
        title="Shivani AI System Overview",
        project_name="Shivani 1.0",
        problem_statement="Fragmented desktop automation and lack of multimodal agency on Windows.",
        solution_summary="A voice-first, multimodal, secure, local-first personal AI operating layer.",
        tech_stack=["Python 3.12", "FastAPI", "Playwright", "pywinauto", "SQLite"],
        key_features=["Voice & Wake Word", "Computer & Browser Control", "Security Sandbox", "Offline Autonomy"],
        mode=PresentationMode.HACKATHON,
        target_duration=5
    )
    
    # 1. Deck metadata validation
    assert deck is not None
    assert deck.total_slides >= 5
    assert Path(deck.pptx_path).exists()
    assert deck.pptx_path.endswith(".pptx")
    
    # 2. Inspect real PPTX file using python-pptx
    prs = Presentation(deck.pptx_path)
    assert len(prs.slides) == deck.total_slides
    
    # Inspect first slide (title slide)
    first_slide = prs.slides[0]
    slide_texts = []
    for shape in first_slide.shapes:
        if shape.has_text_frame:
            for paragraph in shape.text_frame.paragraphs:
                slide_texts.append(paragraph.text)
    joined_text = " ".join(slide_texts)
    assert "Shivani" in joined_text
    
    # 3. Pitch scripts verification
    assert deck.pitch_bundle is not None
    assert "pitch_30s" in deck.pitch_bundle
    assert "pitch_1m" in deck.pitch_bundle
    assert "pitch_3m" in deck.pitch_bundle
    assert "pitch_5m" in deck.pitch_bundle
    assert len(deck.pitch_bundle["pitch_1m"]) > 20
    
    # 4. Judge Q&A bundle verification
    assert deck.qa_bundle is not None
    assert len(deck.qa_bundle) >= 5
    categories = {q["category"] for q in deck.qa_bundle}
    assert len(categories) >= 3
