"""
SHIVANI Multi-Format Document Parser
Parses Markdown, Code, Plaintext, JSON/YAML, CSV/TSV, and PowerPoint presentations
into a structured representation preserving section headers, line boundaries, and tabular data.
"""

from dataclasses import dataclass, field
import json
import os
from pathlib import Path
import re
from typing import Any, Dict, List, Optional, Tuple

from knowledge.ingestion.table_extractor import TableExtractor
from knowledge.security.secret_scanner import SecretScanner


@dataclass
class HeadingNode:
    level: int
    title: str
    line_number: int


@dataclass
class ParsedDocument:
    file_path: str
    file_name: str
    extension: str
    title: str
    content: str
    headings: List[HeadingNode] = field(default_factory=list)
    tables: List[Dict[str, Any]] = field(default_factory=list)
    line_count: int = 1
    metadata: Dict[str, Any] = field(default_factory=dict)


class DocumentParser:
    """Parses various file formats into unified ParsedDocument structures."""

    def __init__(self, redact_secrets: bool = True):
        self.redact_secrets = redact_secrets

    def parse_file(self, file_path: str) -> ParsedDocument:
        """Loads and parses a document from file path."""
        p = Path(file_path).resolve()
        if not p.is_file():
            raise FileNotFoundError(f"File not found: {file_path}")

        ext = p.suffix.lower()

        if ext == ".pptx":
            return self._parse_pptx(p)
        elif ext in (".csv", ".tsv"):
            return self._parse_csv_tsv(p, ext)
        else:
            return self._parse_text_based(p, ext)

    def _parse_text_based(self, path: Path, ext: str) -> ParsedDocument:
        try:
            raw_text = path.read_text(encoding="utf-8", errors="replace")
        except Exception as e:
            raw_text = f"Error reading file {path.name}: {e}"

        # Redact secrets and neutralize prompt injections
        if self.redact_secrets:
            clean_text = SecretScanner.clean_document(raw_text)
        else:
            clean_text = raw_text

        lines = clean_text.splitlines()
        line_count = max(len(lines), 1)

        # Extract headings (supports Markdown #, ##, etc.)
        headings: List[HeadingNode] = []
        heading_re = re.compile(r"^(#{1,6})\s+(.+)$")
        for i, line in enumerate(lines):
            m = heading_re.match(line.strip())
            if m:
                level = len(m.group(1))
                htitle = m.group(2).strip()
                headings.append(HeadingNode(level=level, title=htitle, line_number=i + 1))

        # Extract title: from first H1 or first non-empty line or file name
        title = path.stem.replace("_", " ").title()
        if headings and headings[0].level == 1:
            title = headings[0].title
        elif lines:
            for l in lines:
                s = l.strip()
                if s and not s.startswith(("#", "//", "/*", "'''", '"""')):
                    title = s[:80]
                    break

        # Extract tables if markdown
        tables: List[Dict[str, Any]] = []
        if ext in (".md", ".markdown"):
            tables = TableExtractor.extract_markdown_tables(clean_text)

        metadata: Dict[str, Any] = {
            "size_bytes": path.stat().st_size if path.exists() else len(raw_text.encode("utf-8")),
            "last_modified": path.stat().st_mtime if path.exists() else 0.0,
            "heading_count": len(headings),
        }

        return ParsedDocument(
            file_path=str(path),
            file_name=path.name,
            extension=ext,
            title=title,
            content=clean_text,
            headings=headings,
            tables=tables,
            line_count=line_count,
            metadata=metadata,
        )

    def _parse_csv_tsv(self, path: Path, ext: str) -> ParsedDocument:
        delimiter = "\t" if ext == ".tsv" else ","
        try:
            raw = path.read_text(encoding="utf-8", errors="replace")
        except Exception:
            raw = ""

        if self.redact_secrets:
            raw = SecretScanner.clean_document(raw)

        md_content = TableExtractor.parse_csv_to_markdown(raw, delimiter=delimiter)
        return ParsedDocument(
            file_path=str(path),
            file_name=path.name,
            extension=ext,
            title=path.stem.replace("_", " ").title(),
            content=md_content,
            headings=[HeadingNode(level=1, title=path.stem, line_number=1)],
            tables=[{"markdown": md_content, "line_start": 1, "line_end": len(md_content.splitlines())}],
            line_count=len(md_content.splitlines()),
            metadata={"row_count": len(raw.splitlines())},
        )

    def _parse_pptx(self, path: Path) -> ParsedDocument:
        """Extracts text runs and slide titles from PowerPoint presentations."""
        headings: List[HeadingNode] = []
        slides_text: List[str] = []
        try:
            from pptx import Presentation
            prs = Presentation(str(path))
            cur_line = 1
            for slide_idx, slide in enumerate(prs.slides, start=1):
                slide_title = f"Slide {slide_idx}"
                slide_lines: List[str] = []
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text:
                        text = shape.text.strip()
                        if not slide_lines:
                            slide_title = f"Slide {slide_idx}: {text.splitlines()[0][:50]}"
                        slide_lines.append(text)
                
                heading_node = HeadingNode(level=2, title=slide_title, line_number=cur_line)
                headings.append(heading_node)
                
                block = f"## {slide_title}\n\n" + "\n\n".join(slide_lines)
                slides_text.append(block)
                cur_line += len(block.splitlines()) + 2
                
            full_text = "\n\n---\n\n".join(slides_text)
        except Exception as e:
            full_text = f"PPTX presentation: {path.name} (Error parsing slides: {e})"

        if self.redact_secrets:
            full_text = SecretScanner.clean_document(full_text)

        return ParsedDocument(
            file_path=str(path),
            file_name=path.name,
            extension=".pptx",
            title=path.stem.replace("_", " ").title(),
            content=full_text,
            headings=headings,
            tables=[],
            line_count=len(full_text.splitlines()),
            metadata={"slide_count": len(headings)},
        )
